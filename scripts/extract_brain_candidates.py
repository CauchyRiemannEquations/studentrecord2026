from __future__ import annotations

import fnmatch
import json
import re
from dataclasses import dataclass
from pathlib import Path

import fitz
from pptx import Presentation


ROOT = Path(__file__).resolve().parents[1]
BRAIN_DIR = ROOT / "brain"
REVIEW_DIR = BRAIN_DIR / "review"


@dataclass(frozen=True)
class SourceSpec:
    doc_id: str
    kind: str
    patterns: tuple[str, ...]


SOURCE_SPECS = (
    SourceSpec(
        doc_id="official_guidelines_2026_hs",
        kind="pdf",
        patterns=(
            "서울특별시교육청 학생역량·혁신교육과_2026 학교생활기록부 기재요령(고등학교).pdf",
            "*2026*학교생활기록부*기재요령*고등학교*.pdf",
            "2026 학교생활기록부 기재요령(고).pdf",
        ),
    ),
    SourceSpec(
        doc_id="guide_2026_hs",
        kind="pdf",
        patterns=(
            "2026학년도 학교생활기록부 기재 길라잡이(고등학교).pdf",
            "*2026*학교생활기록부*기재 길라잡이*고등학교*.pdf",
        ),
    ),
    SourceSpec(
        doc_id="changes_ppt_2026_hs",
        kind="pptx",
        patterns=(
            "2026 학교생활기록부 기재요령 주요 개정사항_고등_F_260218.pptx",
            "*2026*학교생활기록부*주요 개정사항*.pptx",
        ),
    ),
    SourceSpec(
        doc_id="training_ppt_2026_hs",
        kind="pptx",
        patterns=(
            "2026 학교생활기록부 연수(성북강북교육지원청).pptx",
            "*2026*학교생활기록부*연수*.pptx",
        ),
    ),
)

GUIDE_QA_RANGES = [
    (92, 93, "processing"),
    (94, 105, "personal_school_status"),
    (106, 122, "attendance"),
    (123, 123, "awards"),
    (124, 124, "school_violence"),
    (125, 138, "creative_experience"),
    (139, 139, "daily_life_activity"),
    (140, 151, "subject_progress"),
    (152, 153, "reading"),
    (154, 155, "behavior"),
    (156, 159, "corrections"),
]

TOPIC_KEYWORDS = {
    "attendance": ["출결", "결석", "지각", "조퇴", "결과", "개근", "과목출석률"],
    "awards": ["수상", "교내상", "대회", "시상"],
    "certificates": ["자격증", "국가공인", "국가기술자격", "NCS"],
    "school_violence": ["학교폭력", "조치사항", "전학", "출석정지"],
    "creative_experience": ["창의적 체험활동", "창체", "자율", "동아리", "진로"],
    "volunteer_service": ["봉사활동", "1365", "VMS", "DOVOL", "헌혈"],
    "subject_progress": ["교과학습발달상황", "세부능력", "세특", "학점", "단위수", "석차등급"],
    "reading": ["독서", "도서명", "독서활동상황"],
    "behavior": ["행동특성", "종합의견", "행특"],
    "corrections": ["정정", "정정대장", "증빙자료", "학업성적관리위원회"],
    "personal_school_status": ["학적", "전입", "전출", "복학", "재입학"],
    "processing": ["영문", "한글", "한자", "처리요령", "유의사항"],
}

RULE_SIGNAL_PATTERNS = {
    "forbidden": [
        r"입력하지 않는다",
        r"기재하지 않는다",
        r"입력할 수 없다",
        r"기재할 수 없다",
        r"금지",
    ],
    "allowed": [
        r"입력할 수 있다",
        r"기재할 수 있다",
        r"인정한다",
    ],
    "school_rule_needed": [
        r"학교장이 정",
        r"학교교육계획",
        r"학칙",
    ],
    "committee_needed": [
        r"학업성적관리위원회",
        r"심의를 통해",
        r"심의 절차",
    ],
}


def candidate_roots() -> list[Path]:
    home = Path.home()
    roots = [
        home / "Downloads",
        home / "Documents" / "에듀파인",
        home / "Documents" / "카카오톡 받은 파일",
        home / "Documents" / "CoolMessenger Files" / "Received Files",
        home / "Documents",
        ROOT,
    ]
    return [root for root in roots if root.exists()]


def clean_text(text: str) -> str:
    text = text.replace("\x0b", "\n").replace("\u00a0", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def flatten_text(text: str) -> str:
    return " ".join(clean_text(text).split())


def detect_printed_page(text: str) -> int | None:
    for line in [line.strip() for line in clean_text(text).splitlines() if line.strip()][:4]:
        if re.fullmatch(r"\d{1,3}", line):
            return int(line)
    return None


def find_source_path(spec: SourceSpec) -> Path:
    roots = candidate_roots()
    for root in roots:
        for pattern in spec.patterns:
            for candidate in root.rglob("*"):
                if candidate.is_file() and fnmatch.fnmatch(candidate.name, pattern):
                    return candidate
    raise FileNotFoundError(f"Could not find source file for {spec.doc_id}")


def write_json(path: Path, payload: object) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )


def write_jsonl(path: Path, records: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    body = "\n".join(json.dumps(record, ensure_ascii=False) for record in records)
    path.write_text(body + ("\n" if body else ""), encoding="utf-8")


def infer_topic(text: str) -> str | None:
    for topic_id, keywords in TOPIC_KEYWORDS.items():
        if any(keyword in text for keyword in keywords):
            return topic_id
    return None


def qa_topic_from_page(page_number: int) -> str | None:
    for start, end, topic_id in GUIDE_QA_RANGES:
        if start <= page_number <= end:
            return topic_id
    return None


def extract_pdf_units(doc_id: str, path: Path) -> list[dict]:
    units: list[dict] = []
    doc = fitz.open(path)
    try:
        for page_number in range(1, doc.page_count + 1):
            raw_text = doc.load_page(page_number - 1).get_text("text")
            text = clean_text(raw_text)
            units.append(
                {
                    "docId": doc_id,
                    "unitType": "page",
                    "page": page_number,
                    "printedPageHint": detect_printed_page(raw_text),
                    "text": text,
                    "flatText": flatten_text(raw_text),
                    "charCount": len(text),
                    "sourcePath": str(path),
                }
            )
    finally:
        doc.close()
    return units


def extract_ppt_units(doc_id: str, path: Path) -> list[dict]:
    units: list[dict] = []
    prs = Presentation(path)
    for slide_number, slide in enumerate(prs.slides, start=1):
        parts: list[str] = []
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                raw_text = getattr(shape, "text", "")
                raw_text = raw_text.replace("\x0b", "\n")
                normalized = clean_text(raw_text)
                if normalized:
                    parts.append(normalized)
        joined = "\n".join(parts)
        units.append(
            {
                "docId": doc_id,
                "unitType": "slide",
                "slide": slide_number,
                "text": joined,
                "flatText": flatten_text(joined),
                "charCount": len(joined),
                "sourcePath": str(path),
            }
        )
    return units


def build_outline_candidates(units: list[dict]) -> list[dict]:
    candidates: list[dict] = []
    interesting = {
        "official_guidelines_2026_hs": {4, 5},
        "guide_2026_hs": {4, 5},
    }
    for unit in units:
        pages = interesting.get(unit["docId"])
        if not pages:
            continue
        page_value = unit.get("page")
        if page_value in pages:
            candidates.append(
                {
                    "candidateType": "page_outline",
                    "docId": unit["docId"],
                    "page": page_value,
                    "text": unit["text"],
                }
            )
    return candidates


def build_change_candidates(units: list[dict]) -> list[dict]:
    candidates: list[dict] = []
    for unit in units:
        if unit["docId"] != "changes_ppt_2026_hs":
            continue
        text = unit["flatText"]
        if "2026" not in text or "2025" not in text:
            continue
        split_match = re.search(r"\b2025\b", text)
        if not split_match:
            continue
        before = text[: split_match.start()].strip(" |")
        after = text[split_match.start() :].strip(" |")
        if before.startswith("2026"):
            before = before[4:].strip(" |")
        topic_guess = infer_topic(text)
        candidates.append(
            {
                "candidateType": "change_compare_slide",
                "candidateId": f"change_slide_{unit['slide']:03d}",
                "docId": unit["docId"],
                "slide": unit["slide"],
                "topicGuess": topic_guess,
                "after2026Raw": before,
                "before2025Raw": after,
                "rawText": text,
                "reviewHint": "슬라이드 비교 문구를 사람이 완성문으로 다듬어 04_change_2026_catalog.jsonl 한 줄로 승격",
            }
        )
    return candidates


def paragraph_snippets(text: str) -> list[str]:
    blocks = re.split(r"\n\s*\n", text)
    normalized = []
    for block in blocks:
        block = clean_text(block)
        if not block:
            continue
        normalized.append(block)
    if normalized:
        return normalized
    return [clean_text(text)] if clean_text(text) else []


def build_rule_candidates(units: list[dict]) -> list[dict]:
    candidates: list[dict] = []
    for unit in units:
        if unit["docId"] not in {"official_guidelines_2026_hs", "guide_2026_hs"}:
            continue
        seen_snippets: set[str] = set()
        count_for_page = 0
        for snippet in paragraph_snippets(unit["text"]):
            signals = [
                label
                for label, patterns in RULE_SIGNAL_PATTERNS.items()
                if any(re.search(pattern, snippet) for pattern in patterns)
            ]
            if not signals:
                continue
            short = flatten_text(snippet)
            if short in seen_snippets:
                continue
            seen_snippets.add(short)
            candidates.append(
                {
                    "candidateType": "rule_snippet",
                    "docId": unit["docId"],
                    "page": unit.get("page"),
                    "printedPageHint": unit.get("printedPageHint"),
                    "topicGuess": infer_topic(short),
                    "signals": signals,
                    "snippet": short,
                    "reviewHint": "허용/금지/조건/학교규정/위원회 여부를 사람이 판별해 rule card로 승격",
                }
            )
            count_for_page += 1
            if count_for_page >= 8:
                break
    return candidates


def extract_tags(text: str) -> list[str]:
    return sorted(set(re.findall(r"#[^#\s]+", text)))


def split_question_answer(text: str) -> tuple[str, str]:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    if "❶" in lines:
        lines = lines[lines.index("❶") + 1 :]
    question_lines: list[str] = []
    answer_lines: list[str] = []
    seen_question = False
    for line in lines:
        if line.startswith("「") or line.startswith("<"):
            break
        if not seen_question:
            question_lines.append(line)
            if "?" in line:
                seen_question = True
            continue
        if line.endswith("?"):
            question_lines.append(line)
            continue
        answer_lines.append(line)
    return (" ".join(question_lines[:5]).strip(), " ".join(answer_lines[:6]).strip())


def build_qa_candidates(units: list[dict]) -> list[dict]:
    candidates: list[dict] = []
    for unit in units:
        if unit["docId"] != "guide_2026_hs":
            continue
        page_number = unit.get("page")
        if not isinstance(page_number, int) or page_number < 92:
            continue
        topic_id = qa_topic_from_page(page_number)
        question_preview, answer_preview = split_question_answer(unit["text"])
        candidates.append(
            {
                "candidateType": "qa_page_candidate",
                "docId": unit["docId"],
                "page": page_number,
                "topicId": topic_id,
                "tags": extract_tags(unit["text"]),
                "questionPreview": question_preview,
                "answerPreview": answer_preview,
                "rawText": unit["text"],
                "continuationLikely": "❶" not in unit["text"],
                "reviewHint": "페이지 단위 후보를 묶어 질문·답변이 함께 있는 QA 카드로 정리",
            }
        )
    return candidates


def build_term_candidates(units: list[dict]) -> list[dict]:
    candidates: list[dict] = []
    for unit in units:
        if unit["docId"] != "guide_2026_hs":
            continue
        tags = extract_tags(unit["text"])
        if not tags:
            continue
        candidates.append(
            {
                "candidateType": "term_tags",
                "docId": unit["docId"],
                "page": unit["page"],
                "topicId": qa_topic_from_page(unit["page"]),
                "tags": tags,
            }
        )
    return candidates


def main() -> None:
    REVIEW_DIR.mkdir(parents=True, exist_ok=True)

    source_paths = {spec.doc_id: find_source_path(spec) for spec in SOURCE_SPECS}
    all_units: list[dict] = []

    for spec in SOURCE_SPECS:
        path = source_paths[spec.doc_id]
        if spec.kind == "pdf":
            units = extract_pdf_units(spec.doc_id, path)
        else:
            units = extract_ppt_units(spec.doc_id, path)
        all_units.extend(units)
        write_jsonl(REVIEW_DIR / f"unit_chunks_{spec.doc_id}.jsonl", units)

    outline_candidates = build_outline_candidates(all_units)
    change_candidates = build_change_candidates(all_units)
    rule_candidates = build_rule_candidates(all_units)
    qa_candidates = build_qa_candidates(all_units)
    term_candidates = build_term_candidates(all_units)

    write_jsonl(REVIEW_DIR / "page_outline_candidates.jsonl", outline_candidates)
    write_jsonl(REVIEW_DIR / "change_candidates.jsonl", change_candidates)
    write_jsonl(REVIEW_DIR / "rule_candidates.jsonl", rule_candidates)
    write_jsonl(REVIEW_DIR / "qa_candidates.jsonl", qa_candidates)
    write_jsonl(REVIEW_DIR / "term_tag_candidates.jsonl", term_candidates)

    manifest = {
        "generatedBy": "scripts/extract_brain_candidates.py",
        "outputDir": str(REVIEW_DIR),
        "sourceFiles": {doc_id: str(path) for doc_id, path in source_paths.items()},
        "counts": {
            "totalUnits": len(all_units),
            "pageOutlineCandidates": len(outline_candidates),
            "changeCandidates": len(change_candidates),
            "ruleCandidates": len(rule_candidates),
            "qaCandidates": len(qa_candidates),
            "termTagCandidates": len(term_candidates),
        },
    }
    write_json(REVIEW_DIR / "extraction_manifest.json", manifest)


if __name__ == "__main__":
    main()
